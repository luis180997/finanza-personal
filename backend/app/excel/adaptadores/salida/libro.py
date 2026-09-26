"""El puerto `LectorDeLibro` con openpyxl."""
from __future__ import annotations

import io
from collections.abc import Iterator


class LibroOpenpyxl:
    def __init__(self, wb):
        self._wb = wb
        self.hojas: list[str] = wb.sheetnames

    def filas(self, hoja: str) -> Iterator[tuple]:
        return self._wb[hoja].iter_rows(values_only=True)


class LectorOpenpyxl:
    def abrir(self, contenido: bytes) -> LibroOpenpyxl:
        from openpyxl import load_workbook

        return LibroOpenpyxl(load_workbook(io.BytesIO(contenido), data_only=True, read_only=True))
